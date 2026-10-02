"""doda.domain.knowledge.text_extraction — real parsing libraries against
real (if minimal) documents, not sentinel bytes. FR-KNW-001's own
fixtures (tests/unit/test_file_validation.py) only need to pass
magic-byte/container checks, not be fully parseable — these tests build
genuinely parseable minimal documents instead, since that's what this
module actually exercises.
"""

import io
import zipfile

import openpyxl
import pytest
from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter

from doda.domain.knowledge.file_validation import DOCX, JPEG, PDF, TXT, XLSX
from doda.domain.knowledge.text_extraction import DocumentIndexingError, extract_text

# A hand-built minimal PDF with a broken xref table — pypdf's default
# strict=False recovery mode (verified directly against this exact
# payload before writing this fixture, not assumed) scans for objects
# instead of trusting the xref, so this is a genuinely parseable,
# one-page PDF containing the literal text "Hello World".
_CONTENT_STREAM = b"BT /F1 24 Tf 10 100 Td (Hello World) Tj ET"
_MINIMAL_PDF_BYTES = (
    b"%PDF-1.4\n"
    b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
    b"3 0 obj\n<< /Type /Page /Parent 2 0 R "
    b"/Resources << /Font << /F1 4 0 R >> >> "
    b"/MediaBox [0 0 200 200] /Contents 5 0 R >>\nendobj\n"
    b"4 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    b"5 0 obj\n<< /Length "
    + str(len(_CONTENT_STREAM)).encode()
    + b" >>\nstream\n"
    + _CONTENT_STREAM
    + b"\nendstream\nendobj\n"
    b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n0\n%%EOF\n"
)


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_txt_is_decoded_directly() -> None:
    assert extract_text(content_type=TXT, data=b"salom dunyo") == "salom dunyo"


def test_txt_with_invalid_utf8_is_replaced_not_raised() -> None:
    # file_validation already requires text/plain uploads to be valid
    # UTF-8, so this is a defensive fallback, not a reachable production
    # path — proven not to raise either way.
    result = extract_text(content_type=TXT, data=b"\xff\xfe broken")
    assert result is not None


def test_a_real_minimal_pdfs_text_is_extracted() -> None:
    assert extract_text(content_type=PDF, data=_MINIMAL_PDF_BYTES) == "Hello World"


def test_a_password_protected_pdf_raises_document_indexing_error() -> None:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_MINIMAL_PDF_BYTES)))
    writer.encrypt(user_password="secret", owner_password="secret2")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentIndexingError, match="password-protected"):
        extract_text(content_type=PDF, data=buffer.getvalue())


def test_a_real_docxs_paragraphs_are_extracted() -> None:
    data = _docx_bytes(["Birinchi paragraf.", "Ikkinchi paragraf."])
    text = extract_text(content_type=DOCX, data=data)
    assert text == "Birinchi paragraf.\nIkkinchi paragraf."


def test_a_docx_whose_document_xml_is_malformed_raises_document_indexing_error() -> None:
    """Passed file_validation's own OOXML-container check (a real
    word/document.xml member exists) but the member's content itself is
    not valid XML — validation cannot see this, only parsing can."""
    good = _docx_bytes(["hello"])
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(good)) as source, zipfile.ZipFile(buffer, "w") as dest:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "word/document.xml":
                content = b"<not-even-well-formed-xml"
            dest.writestr(item, content)

    with pytest.raises(DocumentIndexingError, match="DOCX"):
        extract_text(content_type=DOCX, data=buffer.getvalue())


def test_a_real_xlsxs_cells_are_extracted() -> None:
    data = _xlsx_bytes([["ism", "yosh"], ["Ali", 30], ["Vali", 25]])
    text = extract_text(content_type=XLSX, data=data)
    assert text == "ism yosh\nAli 30\nVali 25"


def test_an_image_has_no_extractor_and_returns_none_not_an_error() -> None:
    assert extract_text(content_type=JPEG, data=b"\xff\xd8\xff fake jpeg bytes") is None


def test_an_unrecognized_content_type_returns_none() -> None:
    assert extract_text(content_type="application/octet-stream", data=b"anything") is None
