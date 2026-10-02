"""FR-KNW-002's "parsing" stage: turns a validated upload's raw bytes
into plain text to chunk and embed. Only the four textual formats
file_validation.py accepts are handled — PNG/JPEG (images) have no
extractor here and return None, meaning "not indexed yet" rather than a
failure: an image upload still succeeds via knowledge_service.ingest_file,
it simply has zero DocumentChunk rows (no retrieval UI exists to notice
the gap yet either). OCR/vision-based image indexing is a real,
deliberately deferred extension, not an oversight.

By the time a caller reaches this module, doda.domain.knowledge.
file_validation has already confirmed the bytes structurally match the
claimed type (magic bytes, OOXML container contents) — so extraction
failing here (DocumentIndexingError) means something validation cannot
see: a password-protected PDF, a corrupted-but-structurally-valid-ZIP
DOCX/XLSX, or similar. This is deliberately a DIFFERENT error from
file_validation's own FileValidationError family — a file that failed
validation never reaches this module at all, while a file that failed
extraction already passed validation and was already stored.
"""

import io

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from doda.domain.knowledge.file_validation import DOCX, PDF, TXT, XLSX


class DocumentIndexingError(Exception):
    """Raised when a file that already passed file_validation's
    structural checks still cannot be parsed into text — see module
    docstring for why this is distinct from FileValidationError."""


def extract_text(*, content_type: str, data: bytes) -> str | None:
    """Returns None for a content type with no extractor (not indexed
    yet, not a failure) or DocumentIndexingError if extraction itself
    fails for a type that IS supposed to be extractable."""
    if content_type == TXT:
        return data.decode("utf-8", errors="replace")
    if content_type == PDF:
        return _extract_pdf(data)
    if content_type == DOCX:
        return _extract_docx(data)
    if content_type == XLSX:
        return _extract_xlsx(data)
    return None


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentIndexingError("PDF is password-protected and cannot be parsed")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentIndexingError:
        raise
    except Exception as exc:
        raise DocumentIndexingError(f"failed to parse PDF: {type(exc).__name__}") from None


def _extract_docx(data: bytes) -> str:
    try:
        document = DocxDocument(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception as exc:
        raise DocumentIndexingError(f"failed to parse DOCX: {type(exc).__name__}") from None


def _extract_xlsx(data: bytes) -> str:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None]
                    if cells:
                        lines.append(" ".join(cells))
            return "\n".join(lines)
        finally:
            workbook.close()
    except Exception as exc:
        raise DocumentIndexingError(f"failed to parse XLSX: {type(exc).__name__}") from None
